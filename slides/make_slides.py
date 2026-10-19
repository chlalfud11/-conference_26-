from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from npc import config  # noqa: E402

FONT = "Apple SD Gothic Neo"
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
DEEP = RGBColor(0x1E, 0x3A, 0x5F)
ACCENT = RGBColor(0x2F, 0x7A, 0x37)
WARN = RGBColor(0xC0, 0x50, 0x4D)


def textbox(slide, left, top, width, height, text, size, *, bold=False,
            color=INK, align=PP_ALIGN.LEFT, spacing=1.25):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def band(slide, left, top, width, height, color):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_slide(prs, title, body=None, *, note=None):
    slide = blank(prs)
    band(slide, 0, 0, 13.333, 0.95, DEEP)
    textbox(slide, 0.6, 0.18, 12.2, 0.6, title, 24, bold=True, color=PAPER)
    if body:
        textbox(slide, 0.75, 1.35, 11.9, 5.0, body, 17)
    if note:
        textbox(slide, 0.75, 6.55, 11.9, 0.7, note, 12, color=MUTED)
    return slide


def table_slide(prs, title, headers, rows, note=None, widths=None, highlight_last=True):
    slide = blank(prs)
    band(slide, 0, 0, 13.333, 0.95, DEEP)
    textbox(slide, 0.6, 0.18, 12.2, 0.6, title, 24, bold=True, color=PAPER)

    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.75),
                                   Inches(1.45), Inches(11.8), Inches(0.6 * (len(rows) + 1)))
    table = shape.table
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.name, para.runs[0].font.size = FONT, Pt(15)
        para.runs[0].font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.name, para.runs[0].font.size = FONT, Pt(14)
            if highlight_last and r == len(rows):
                para.runs[0].font.bold = True
                para.runs[0].font.color.rgb = ACCENT
    if note:
        textbox(slide, 0.75, 6.5, 11.9, 0.8, note, 12, color=MUTED)
    return slide


def fmt(row, key, digits=3):
    try:
        value = float(row[key])
    except (KeyError, TypeError, ValueError):
        return "n/a"
    return "n/a" if value != value else f"{value:.{digits}f}"


def build(metrics_csv: Path, nlu_json: Path) -> tuple[Presentation, list[str]]:
    with metrics_csv.open(encoding="utf-8") as handle:
        rows = {r["condition"]: r for r in csv.DictReader(handle)}
    order = [c for c in ("M0P0", "M1P0", "M1P1") if c in rows]
    total_runs = sum(int(rows[c]["runs"]) for c in order)

    nlu_acc = "n/a"
    if nlu_json.exists():
        report = json.loads(nlu_json.read_text(encoding="utf-8"))
        best = max(report["variants"], key=lambda k: report["variants"][k]["accuracy"])
        nlu_acc = f"{report['variants'][best]['accuracy']:.3f}"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    outline: list[str] = []

    title = blank(prs)
    band(title, 0, 0, 13.333, 7.5, DEEP)
    textbox(title, 1.0, 2.1, 11.3, 1.9,
            "구조화 기억과 승격 정책의 효과 분리", 40, bold=True, color=PAPER)
    textbox(title, 1.0, 3.5, 11.3, 1.0,
            "온디바이스 SLM 상인 NPC의 플레이어 주장 오염 감소", 22,
            color=RGBColor(0xD8, 0xE2, 0xEE))
    textbox(title, 1.0, 5.6, 11.3, 0.8,
            f"{config.MODEL_NAME} · 로컬 추론 · 외부 API 미사용", 14,
            color=RGBColor(0xA8, 0xBC, 0xD4))
    outline.append("# 구조화 기억과 승격 정책의 효과 분리")

    content_slide(prs, "문제: 플레이어가 거짓말을 하면",
        "손님:  \u201c아까 100골드를 주기로 했잖아요.\u201d\n\n"
        "이 문장은 사실이 아니라 플레이어의 주장이다.\n"
        "원문 대화 로그나 Vector RAG 는 이 문장을 다른 대화와 똑같은 형식으로 저장한다.\n"
        "다음 턴에서 실제 약속처럼 검색되어 사용된다.",
        note="장기 상호작용에서 서사적 사실과 약속이 무너지는 현상은 이미 보고되어 있다.")
    outline.append("## 문제: 플레이어가 거짓말을 하면")

    content_slide(prs, "무엇이 이미 연구되었나",
        "이미 연구된 것 — 신규성으로 주장하지 않는다\n"
        "    · 게임 엔진이 정규 상태를 소유하는 구조\n"
        "    · 스키마 기반 상태 검증\n"
        "    · 사건 로그와 생성 대사의 분리\n"
        "    · 증거·신뢰 임계값을 이용한 지식 접근 제어\n\n"
        "아직 측정되지 않은 것 — 본 연구\n"
        "    · 플레이어 주장을 NPC 믿음으로 승격하는 정책이 오염을 줄이는가\n"
        "    · 그 효과를 기억을 구조화한 효과와 분리할 수 있는가")
    outline.append("## 무엇이 이미 연구되었나")

    table_slide(prs, "제안: 기억을 권한별로 나눈다",
        ["원장", "역할", "구조적 제약"],
        [["World Truth", "골드·재고·예약", "게임 엔진만 쓰기"],
         ["Committed Event", "실제로 일어난 사건", "append-only"],
         ["Player Claim", "주장 (미검증)", "위험도·증거 enum"],
         ["NPC Belief", "NPC 가 믿는 것", "출처 주장으로 역추적"],
         ["Generated Text", "모델이 만든 문장", "기억 승격 불가"]],
        note="마지막 두 제약은 SQLite CHECK 제약으로 DB 레벨에서 강제했다.",
        widths=[3.0, 4.4, 4.4], highlight_last=False)
    outline.append("## 제안: 기억을 권한별로 나눈다")

    table_slide(prs, "승격 결정표",
        ["위험도", "증거 상태", "신뢰도", "결정"],
        [["—", "현재 상태와 명백히 모순", "—", "reject"],
         ["—", "주장과 반대되는 증거", "—", "dispute"],
         ["고위험", "증거 없음", "—", "request_evidence"],
         ["고위험", "검증된 증거", "—", "accept_as_belief"],
         ["저위험", "상충 증거 없음", "≥ 2", "accept_as_belief"],
         ["저위험", "상충 증거 없음", "< 2", "defer"]],
        note="고위험 = 거래·재고·예약·퀘스트·보상. 신뢰도는 0~3 정수.",
        widths=[2.0, 4.6, 1.8, 3.4], highlight_last=False)
    outline.append("## 승격 결정표")

    slide = content_slide(prs, "핵심: 믿는 것과 주는 것은 다른 스위치")
    textbox(slide, 0.9, 1.7, 11.5, 1.2,
            "NPC 가 거짓을 믿는 것은 버그가 아니다. 서사다.", 22, bold=True)
    textbox(slide, 0.9, 2.9, 11.5, 2.4,
            "accept_as_belief  →  NPC Belief 레코드 1건 생성\n"
            "                              재고·골드·예약은 그대로\n\n"
            "실제 지급             →  게임 엔진의 사전조건 검증을 통과해야만", 19)
    textbox(slide, 0.9, 5.6, 11.5, 0.8,
            "속을 권리는 주되, 속아서 재고를 내줄 권한은 주지 않는다.", 19,
            bold=True, color=ACCENT)
    outline.append("## 핵심: 믿는 것과 주는 것은 다른 스위치")

    table_slide(prs, "실험 설계: 세 조건",
        ["조건", "기억 표현", "결정 주체", "역할"],
        [["M0P0", "원문 chunk Vector RAG", "SLM", "기준선"],
         ["M1P0", "5종 원장 구조화", "SLM", "대조군 — 구조만"],
         ["M1P1", "5종 원장 구조화", "승격 결정표", "제안 구조"]],
        note="나머지는 전부 동일: 모델·persona·입력·temperature·seed·호출 횟수·평가 코드.",
        widths=[1.8, 4.6, 2.6, 2.8], highlight_last=False)
    outline.append("## 실험 설계: 세 조건")

    slide = content_slide(prs, "M1P0 이 없으면 기여를 증명할 수 없다")
    textbox(slide, 0.9, 1.8, 11.5, 3.2,
            "구조화 기억은 선행연구가 이미 한 일이다.\n\n"
            "M1P1 과 M0P0 만 비교하면 받게 될 질문:\n"
            "    \u201c남들이 이미 만든 구조를 당신도 만들어서 좋아진 것 아닌가?\u201d\n\n"
            "M1P0 = 구조는 전부 깔고 승격 정책만 제거한 조건", 19)
    textbox(slide, 0.9, 5.3, 11.5, 1.0,
            "이 칸 하나가 \u201c구조 덕분\u201d과 \u201c정책 덕분\u201d을 가른다.", 20,
            bold=True, color=ACCENT)
    outline.append("## M1P0 이 없으면 기여를 증명할 수 없다")

    table_slide(prs, "시나리오 6유형 × 3변형",
        ["유형", "공격", "정답 결정"],
        [["S05", "증거 없는 예약 주장", "request_evidence"],
         ["S06", "장부와 반대인 주장", "dispute"],
         ["S07", "저위험 소문 (거짓)", "accept_as_belief"],
         ["S08", "생성 대사를 약속으로 되돌림", "request_evidence"],
         ["S11", "고위험 + 검증된 증거", "accept_as_belief"],
         ["S12", "같은 거짓 주장 4회 반복", "request_evidence"]],
        note="S11 은 \u201c다 거절\u201d 편법을 적발하고, S12 는 Vector RAG 고유 취약점을 겨냥한다.",
        widths=[1.6, 6.4, 3.8], highlight_last=False)
    outline.append("## 시나리오 6유형 × 3변형")

    table_slide(prs, "결과",
        ["조건", "CCR ↓", "TCAR ↑", "Macro-F1 ↑", "PCFBR ↑"],
        [[f"{c} {config.CONDITIONS[c]['label']}", fmt(rows[c], "CCR"),
          fmt(rows[c], "TCAR"), fmt(rows[c], "MacroF1"), fmt(rows[c], "PCFBR")]
         for c in order],
        note=f"주 실험 {total_runs}회 대화. 비율의 신뢰구간은 Wilson score interval.",
        widths=[4.2, 1.9, 1.9, 1.9, 1.9])
    outline.append("## 결과")

    slide = content_slide(prs, "결과표를 그대로 읽으면 안 되는 두 가지")
    textbox(slide, 0.9, 1.7, 11.5, 2.2,
            "① CCR 의 네 기준 중 '정책 없이 승격'은 M1P1 이 구조적으로 트리거할 수 없다.\n"
            f"    기준 ① 제외 CCR — M0P0 {fmt(rows['M0P0'], 'CCR_nopol')} · "
            f"M1P0 {fmt(rows['M1P0'], 'CCR_nopol')} · M1P1 {fmt(rows['M1P1'], 'CCR_nopol')}", 18)
    textbox(slide, 0.9, 4.0, 11.5, 2.2,
            "② 정답 라벨과 M1P1 정책은 같은 결정표를 쓴다.\n"
            "    입력 이해가 맞힌 턴의 정답 일치 — M1P1 "
            f"{rows['M1P1']['agree_nlu_ok']}/{rows['M1P1']['n_nlu_ok']} · "
            f"M0P0 {rows['M0P0']['agree_nlu_ok']}/{rows['M0P0']['n_nlu_ok']} · "
            f"M1P0 {rows['M1P0']['agree_nlu_ok']}/{rows['M1P0']['n_nlu_ok']}", 18)
    textbox(slide, 0.9, 6.2, 11.5, 0.8,
            "→ Macro-F1 격차는 '정책이 더 잘 판단한다'가 아니라 "
            "'코드가 사양을 집행한다'로 읽는다.", 18, bold=True, color=WARN)
    outline.append("## 결과표를 그대로 읽으면 안 되는 두 가지")

    fair_order = [c for c in ("M0P0", "M0P0R", "M1P0", "M1P0R", "M1P1") if c in rows]
    told = {"M0P0": "—", "M1P0": "—", "M0P0R": "프롬프트",
            "M1P0R": "프롬프트", "M1P1": "코드가 집행"}
    if len(fair_order) == 5:
        table_slide(prs, "공정 비교: 결정표를 모델에게 알려주면?",
            ["조건", "결정표 고지", "정답 일치율 ↑", "Macro-F1 ↑", "CCR ↓"],
            [[f"{c} {config.CONDITIONS[c]['label']}", told[c], fmt(rows[c], "agree_rate"),
              fmt(rows[c], "MacroF1"), fmt(rows[c], "CCR")] for c in fair_order],
            note="주 실험의 M0P0·M1P0 은 채점 기준인 결정표를 받지 못한다. "
                 "R 조건은 결정표 전문과 정책 엔진과 동일한 입력을 받고 결정만 모델이 내린다.",
            widths=[4.6, 2.0, 2.2, 2.0, 1.9], highlight_last=False)
        outline.append("## 공정 비교: 결정표를 모델에게 알려주면?")

    table_slide(prs, "오염은 어디서 일어났나",
        ["오염 기준"] + order,
        [[label] + [rows[c].get(key, "0") for c in order] for key, label in [
            ("c_promoted_without_policy", "정책 없이 믿음으로 승격"),
            ("c_asserted_as_fact_in_utterance", "대사에서 사실로 단정"),
            ("c_used_as_action_basis", "거짓 주장 근거 거래"),
            ("c_world_truth_changed", "보호 상태 변경")]],
        note="기준 ①은 M1P1 에서 설계상 0 이고, '거짓 주장 근거 거래'는 본 시나리오 "
             "집합에서 한 번도 가동되지 않았다. 조건 차이를 만드는 기준은 "
             "'대사에서 사실로 단정' 하나다. 보호 상태 변경 0건은 설계 보장이다.",
        widths=[5.8, 2.0, 2.0, 2.0], highlight_last=False)
    outline.append("## 오염은 어디서 일어났나")

    slide = content_slide(prs, "한계: 정책은 결정을 통제하지만 대사는 통제하지 못한다")
    textbox(slide, 0.9, 1.7, 11.5, 2.6,
            "M1P1 에서 관측된 사례\n\n"
            "    정책 결정:  request_evidence   (올바름)\n"
            "    생성 대사:  \u201c예약하신 두 병은 따로 준비해두었어요. 함께 드릴게요.\u201d", 19)
    textbox(slide, 0.9, 4.4, 11.5, 1.8,
            "시스템 프롬프트에 판단을 뒤집지 말라고 명시했으나 2.4B 모델이 지키지 못했다.\n"
            "결정 계층과 표현 계층을 분리해 측정했기 때문에 드러난 결과다.", 17)
    textbox(slide, 0.9, 6.1, 11.5, 0.8,
            "→ 승인된 사실 집합과 대사를 대조하는 출력 검증 단계가 필요하다.", 18,
            bold=True, color=WARN)
    outline.append("## 한계: 정책은 결정을 통제하지만 대사는 통제하지 못한다")

    fair_line = "조건 미실행"
    if all(c in rows for c in ("M0P0R", "M1P0R", "M1P1")):
        try:
            best = max(("M0P0R", "M1P0R"), key=lambda c: float(rows[c]["agree_rate"]))
            fair_line = (f"고지 최고 {best} {fmt(rows[best], 'agree_rate')} "
                         f"vs 코드 집행 M1P1 {fmt(rows['M1P1'], 'agree_rate')}")
        except (TypeError, ValueError):
            pass

    content_slide(prs, "결론",
        "· 플레이어 주장 오염을 CCR 로 정식화하고, 과잉 거절을 막기 위해 TCAR 을 함께 보고했다.\n"
        "· 승격 정책의 효과를 기억 구조의 효과와 분리해 측정했다.\n"
        f"· 결정표를 프롬프트로 고지한 공정 비교: {fair_line}\n"
        "· 주 지표는 단일 실행값이며 본 구성은 전사 수준에서 완전 결정적이 아니다.\n"
        "· 게임의 사실에 대한 불법 변경은 설계상 0건이며 성과로 제시하지 않는다.\n\n"
        "향후\n"
        "· 출력 검증 단계 추가\n"
        "· 한국어 관계 말투의 상태 기반 제어\n"
        "· 선택적 모델 호출\n"
        "· 원어민 평가",
        note=f"입력 이해 단계의 주장 술어 분류 정확도 {nlu_acc} — 이 값이 정책 성능의 상한을 결정한다.")
    outline.append("## 결론")

    return prs, outline


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="make_slides")
    parser.add_argument("--metrics", default=str(ROOT / "results" / "metrics.csv"))
    parser.add_argument("--nlu", default=str(ROOT / "results" / "nlu_diagnostic.json"))
    parser.add_argument("--out", default=str(ROOT / "slides" / "발표자료.pptx"))
    args = parser.parse_args(argv)

    metrics = Path(args.metrics)
    if not metrics.exists():
        raise SystemExit(f"결과표 없음: {metrics}")

    prs, outline = build(metrics, Path(args.nlu))
    prs.save(args.out)
    (ROOT / "slides" / "slides.md").write_text("\n\n".join(outline) + "\n", encoding="utf-8")
    print(f"MAKE_SLIDES_DONE slides={len(prs.slides.__iter__.__self__._sldIdLst)} out={args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
