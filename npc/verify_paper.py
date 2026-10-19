from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

from . import config

NUMBER = re.compile(r"\d+\.\d{2,4}")


def measured_values(metrics_csv: Path) -> set[str]:
    values: set[str] = set()
    with metrics_csv.open(encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            for key, raw in row.items():
                if key in ("condition", "label"):
                    continue
                try:
                    number = float(raw)
                except (TypeError, ValueError):
                    continue
                values.add(f"{number:.4f}")
                values.add(f"{number:.3f}")
                values.add(f"{number:.2f}")
                values.add(str(int(number)) if number.is_integer() else "")
                values.add(f"{number * 100:.1f}")
                values.add(f"{number * 100:.0f}")
    values.discard("")
    return values


def declared_constants() -> set[str]:
    values = {f"{config.TEMPERATURE}", f"{config.TOP_P}", "1.96", "0.95", "0.65"}
    for seed in config.SEEDS:
        values.add(str(seed))
    return values


def injected_values(manifest: Path) -> set[str]:
    if not manifest.exists():
        return set()
    values: set[str] = set()
    for raw in json.loads(manifest.read_text(encoding="utf-8")).values():
        values.update(NUMBER.findall(str(raw)))
    return values


def document_lines(document: Path) -> list[str]:
    if document.suffix != ".pptx":
        return document.read_text(encoding="utf-8").splitlines()
    from pptx import Presentation

    lines: list[str] = []
    for slide in Presentation(str(document)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(shape.text_frame.text.splitlines())
            if shape.has_table:
                for row in shape.table.rows:
                    lines.extend(cell.text for cell in row.cells)
    return lines


def check(document: Path, allowed: set[str]) -> list[tuple[int, str, str]]:
    offenders = []
    for lineno, line in enumerate(document_lines(document), 1):
        stripped = line.lstrip()
        if stripped.startswith("## \ucc38\uace0\ubb38\ud5cc"):
            break
        if stripped.startswith(("|---", "```", "[")) or "arXiv" in line:
            continue
        for token in NUMBER.findall(line):
            if token not in allowed:
                offenders.append((lineno, token, line.strip()[:90]))
    return offenders


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="npc.verify_paper")
    parser.add_argument("--metrics", default=str(config.ROOT / "results" / "metrics.csv"))
    parser.add_argument("--manifest",
                        default=str(config.ROOT / "results" / "paper_values.json"))
    parser.add_argument("--documents", nargs="+",
                        default=[str(config.ROOT / "paper" / "draft.md"),
                                 str(config.ROOT / "slides" / "slides.md"),
                                 str(config.ROOT / "slides" / "발표자료.pptx")])
    args = parser.parse_args(argv)

    metrics_path = Path(args.metrics)
    if not metrics_path.exists():
        raise SystemExit(f"결과표 없음: {metrics_path}")
    manifest = Path(args.manifest)
    allowed = (measured_values(metrics_path) | declared_constants()
               | injected_values(manifest))

    total = 0
    for name in args.documents:
        document = Path(name)
        if not document.exists():
            print(f"SKIP {document} (없음)")
            continue
        offenders = check(document, allowed)
        total += len(offenders)
        status = "OK" if not offenders else f"MISMATCH {len(offenders)}"
        print(f"{status:<14} {document}")
        for lineno, token, text in offenders:
            print(f"   L{lineno} 미측정 수치 {token!r}: {text}")

    print(f"VERIFY_PAPER {'PASS' if total == 0 else 'FAIL'} offenders={total} "
          f"allowed_values={len(allowed)}")
    return 0 if total == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
